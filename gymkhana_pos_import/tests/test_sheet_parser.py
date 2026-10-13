"""Spreadsheet tests: the same report as .xls and .xlsx.

The sheets are built here rather than committed, so the layout the reader
depends on is written down in one readable place and no club sales data has
to live in the repository.

They need no database. Inside Odoo they run with the module's tests
(``--test-tags /gymkhana_pos_import``); outside Odoo run them with
``python -m unittest discover -s gymkhana_pos_import/tests -p 'test_sheet_parser.py'``.
"""
import io
import sys
from datetime import date
from decimal import Decimal
from pathlib import Path

HERE = Path(__file__).resolve().parent
PDF_FIXTURE = HERE / "fixtures" / "sales_register_2026-09-27.pdf"

try:
    from odoo.tests import BaseCase as TestCase, tagged
    from ..lib import gymkhana_parser as gp
except ImportError:  # standalone run, without Odoo
    import unittest
    TestCase = unittest.TestCase

    def tagged(*_tags):
        return lambda cls: cls
    sys.path[:0] = [str(HERE.parent / "lib")]
    import gymkhana_parser as gp

D = Decimal

PERIOD = " Date From: 01/10/2026 to 01/10/2026  Outlet:Dhostana Ventures Ltd  Ignore Item Sales Rate  "
PRINTED = "02/10/2026 19:15:24"

# The export, column index by column index, exactly as the POS lays it out:
# page furniture, one heading row, then Group / Sub Group / items / totals
# with a blank row between each. Money arrives as numbers, while Item Count
# and Rate arrive as text -- including a thousands separator.
SHEET = [
    {0: "Printed By :", 2: "Dhostana Supervisor", 8: "Nairobi Gymkhana"},
    {2: PRINTED},
    {14: "Page :", 15: 1.0},
    {0: "Printed Date :", 7: "Group Sales Register -[ Group,Sub Group And Item Wise Sales Register [Summary] ]"},
    {}, {}, {}, {},
    {0: PERIOD},
    {},
    {0: "Menu Item", 4: "Unit", 5: "Item Count", 8: "Rate", 10: "Quantity",
     11: "Amount", 12: "Disc. Amt.", 13: "Total Tax", 14: "Net. Amt."},
    {},
    {0: "Group", 2: "Banda Bar"},
    {},
    {0: "Sub Group", 2: "Beer/RTD/0.0/Cans(BB)"},
    {},
    {0: "Tusker Beer 500Ml(BB)", 4: "Bottle", 5: "2", 8: "211.86", 10: 2.0,
     11: 423.72, 12: 0.0, 13: 76.28, 14: 500.0},
    {},
    {0: "White Cap 500Ml(BB)", 4: "Bottle", 5: "12", 8: "241.53", 10: 15.0,
     11: 3622.92, 12: 0.0, 13: 652.08, 14: 4275.0},
    {},
    {0: "Sub Group Total:", 11: 4046.64, 12: 0.0, 13: 728.36, 14: 4775.0},
    {},
    {0: "Sub Group", 2: "Blended Whiskey(BB)"},
    {},
    {0: "Black Label(BB)", 4: "375Ml", 5: "1", 8: "2,245.76", 10: 1.0,
     11: 2245.76, 12: 0.0, 13: 404.24, 14: 2650.0},
    {},
    {0: "Sub Group Total:", 11: 2245.76, 12: 0.0, 13: 404.24, 14: 2650.0},
    {},
    {0: "Group Total:", 11: 6292.40, 12: 0.0, 13: 1132.60, 14: 7425.0},
    {},
    {0: "Grand Total:", 11: 6292.40, 12: 0.0, 13: 1132.60, 14: 7425.0},
]


def write_xls(rows):
    """``rows`` as a legacy .xls workbook, the way the POS exports one."""
    import xlwt

    book = xlwt.Workbook()
    sheet = book.add_sheet("Sheet1")
    for r, cells in enumerate(rows):
        for c, value in cells.items():
            sheet.write(r, c, value)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def write_xlsx(rows):
    """The same rows as a modern .xlsx workbook."""
    import openpyxl

    book = openpyxl.Workbook()
    sheet = book.active
    for r, cells in enumerate(rows):
        for c, value in cells.items():
            sheet.cell(row=r + 1, column=c + 1, value=value)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def edited(**cells_by_row):
    """A copy of SHEET with some cells replaced: ``edited(**{"28": {11: 1.0}})``."""
    rows = [dict(row) for row in SHEET]
    for row_index, cells in cells_by_row.items():
        rows[int(row_index)].update(cells)
    return rows


@tagged("post_install", "-at_install", "gymkhana_pos_import")
class TestSheetParser(TestCase):
    """One report, both spreadsheet formats."""

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.xls = write_xls(SHEET)
        cls.xlsx = write_xlsx(SHEET)
        cls.res_xls = gp.parse(cls.xls)
        cls.res_xlsx = gp.parse(cls.xlsx)

    # -- what kind of file is this ------------------------------------------

    def test_sniff_tells_the_three_kinds_apart(self):
        self.assertEqual(gp.sniff(PDF_FIXTURE.read_bytes()), "pdf")
        self.assertEqual(gp.sniff(self.xls), "xls")
        self.assertEqual(gp.sniff(self.xlsx), "xlsx")

    def test_sniff_reads_the_file_not_the_name(self):
        """A path is opened and looked at, whatever it is called."""
        self.assertEqual(gp.sniff(PDF_FIXTURE), "pdf")

    def test_a_file_that_is_none_of_the_three_says_so(self):
        with self.assertRaises(gp.ParseError) as caught:
            gp.parse(b"just some text, not a report at all")
        self.assertIn("not a PDF", str(caught.exception))

    def test_a_web_page_saved_as_xls_says_so(self):
        """Some exporters hand out HTML named .xls; the message should say that
        rather than complain about a broken workbook."""
        with self.assertRaises(gp.ParseError) as caught:
            gp.parse(b"<html><body><table><tr><td>Menu Item</td></tr></table></body></html>")
        self.assertIn("web page", str(caught.exception))

    # -- the period and the outlet ------------------------------------------

    def test_the_period_is_read_day_first(self):
        """01/10/2026 in a spreadsheet is 1 October, not 10 January."""
        for res in (self.res_xls, self.res_xlsx):
            self.assertEqual(res["header"]["date_from"], date(2026, 10, 1))
            self.assertEqual(res["header"]["date_to"], date(2026, 10, 1))
            self.assertEqual(res["header"]["outlet"], "Dhostana Ventures Ltd")

    def test_the_pdf_is_still_read_month_first(self):
        """The two exports disagree, and changing one must not change the other."""
        res = gp.parse(PDF_FIXTURE.read_bytes())
        self.assertEqual(res["header"]["date_from"], date(2026, 9, 27))

    def test_a_report_printed_before_its_own_period_is_refused(self):
        """What a period read back to front looks like: the stamp lands before
        the day being reported."""
        with self.assertRaises(gp.ParseError) as caught:
            gp.parse(write_xls(edited(**{"1": {2: "30/09/2026 19:15:24"}})))
        self.assertIn("before the day it reports", str(caught.exception))

    def test_a_period_that_ends_before_it_starts_is_refused(self):
        with self.assertRaises(gp.ParseError) as caught:
            gp.parse(write_xls(edited(**{"8": {
                0: " Date From: 05/10/2026 to 01/10/2026  Outlet:Dhostana Ventures Ltd  "}})))
        self.assertIn("ends before it starts", str(caught.exception))

    # -- the lines ----------------------------------------------------------

    def test_both_formats_read_the_same_report(self):
        self.assertEqual(
            [(ln["key"], ln["qty"], ln["net"]) for ln in self.res_xls["lines"]],
            [(ln["key"], ln["qty"], ln["net"]) for ln in self.res_xlsx["lines"]],
        )

    def test_lines_carry_name_unit_bar_and_money(self):
        lines = self.res_xls["lines"]
        self.assertEqual(len(lines), 3)
        first = lines[0]
        self.assertEqual(first["name"], "Tusker Beer 500Ml")
        self.assertEqual(first["unit"], "Bottle")
        self.assertEqual(first["bar_suffix"], "BB")
        self.assertEqual(first["group"], "Banda Bar")
        self.assertEqual(first["sub_group"], "Beer/RTD/0.0/Cans(BB)")
        self.assertEqual(first["qty"], D("2"))
        self.assertEqual(first["net"], D("500"))
        self.assertEqual(first["price_unit"], D("250.00"))
        self.assertTrue(first["price_exact"])

    def test_a_rate_with_a_thousands_separator_is_read(self):
        black_label = self.res_xls["lines"][2]
        self.assertEqual(black_label["rate"], D("2245.76"))
        self.assertEqual(black_label["price_unit"], D("2650.00"))
        self.assertFalse(black_label["rate_mismatch"])

    def test_whole_numbers_survive_the_sheet(self):
        """Money comes out of a sheet as 7425.0, not '7,425.00'. It still has
        to add up to the cent: the net column is checked with no tolerance."""
        total = self.res_xls["grand_total"]
        self.assertEqual(total["net"], D("7425"))
        self.assertEqual(total["amount"], D("6292.40"))
        self.assertEqual(sum(ln["net"] for ln in self.res_xls["lines"]), total["net"])

    def test_the_whole_sheet_is_one_page(self):
        self.assertEqual(self.res_xls["page_count"], 1)

    # -- the checks still bite ----------------------------------------------

    def test_a_grand_total_that_does_not_match_is_refused(self):
        with self.assertRaises(gp.ParseError) as caught:
            gp.parse(write_xls(edited(**{"30": {14: 7420.0}})))
        self.assertIn("7,425.00", str(caught.exception))

    def test_a_sub_group_total_that_does_not_match_is_refused(self):
        with self.assertRaises(gp.ParseError) as caught:
            gp.parse(write_xls(edited(**{"20": {14: 4770.0}})))
        self.assertIn("Beer/RTD/0.0/Cans(BB)", str(caught.exception))

    def test_a_line_that_does_not_add_up_is_refused(self):
        """Amount less discount plus tax must give the net."""
        with self.assertRaises(gp.ParseError) as caught:
            gp.parse(write_xls(edited(**{"16": {13: 70.0}})))
        self.assertIn("doesn't give the net", str(caught.exception))

    def test_a_refund_line_is_refused(self):
        with self.assertRaises(gp.ParseError) as caught:
            gp.parse(write_xls(edited(**{"16": {10: -2.0}})))
        self.assertIn("refunds and voids", str(caught.exception))

    def test_a_sheet_without_the_heading_row_is_refused(self):
        with self.assertRaises(gp.ParseError) as caught:
            gp.parse(write_xls(edited(**{"10": {5: "Count"}})))
        self.assertIn("column header row was not found", str(caught.exception))

    def test_a_sheet_without_the_period_line_is_refused(self):
        rows = [dict(row) for row in SHEET]
        rows[8] = {}
        with self.assertRaises(gp.ParseError) as caught:
            gp.parse(write_xls(rows))
        self.assertIn("'Date From ... Outlet' line was not found", str(caught.exception))

    def test_an_empty_sheet_is_refused(self):
        with self.assertRaises(gp.ParseError) as caught:
            gp.parse(write_xlsx([{}, {}]))
        self.assertIn("sheet is empty", str(caught.exception))
