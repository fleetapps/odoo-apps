"""PDF and Excel: the figures of the screen, computed again server-side."""

import io
from unittest.mock import patch

import openpyxl

from odoo.tests import tagged

from ..models import odin_pnl_report
from .common import PnlCase


@tagged("post_install", "-at_install")
class TestExport(PnlCase):

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.post("2026-09-10", [(cls.revenue, -1000.0, {"partner": cls.partner_a}), (cls.cos, 400.0)])

    def test_the_workbook_holds_real_numbers_for_every_line(self):
        options = self.options(unfolded=["L:REV"], comparison={"mode": "previous_period", "periods": 1})
        content, filename, mimetype = self.env["odin.pnl.export"].render(options, "xlsx")
        self.assertTrue(filename.endswith(".xlsx"))
        sheet = openpyxl.load_workbook(io.BytesIO(content)).active
        values = {row[0].value: [cell.value for cell in row[1:]] for row in sheet.iter_rows(min_row=6)}
        net = next(value for key, value in values.items() if key and key.endswith("Net Profit"))
        self.assertEqual(net[0], 600.0)
        account = next(key for key in values if key and self.revenue.name in key)
        self.assertEqual(values[account][0], 1000.0, "unfolded accounts are exported")

    def test_a_ledger_exports_its_running_balance(self):
        self.post("2026-09-20", [(self.revenue, -250.0)])
        account_key = f"L:REV/A:{self.revenue.id}"
        options = self.options(unfolded=["L:REV", account_key], modes={account_key: "ledger"})
        payload = self.env["odin.pnl.export"]._payload(options)
        labels = [column["label"] for column in payload["columns"]]
        self.assertIn("Balance", labels)
        position = labels.index("Balance")
        entries = [row for row in payload["rows"] if row["type"] == "entry"]
        self.assertEqual([row["cells"][position]["value"] for row in entries], [1000.0, 1250.0])
        account = next(row for row in payload["rows"] if row["key"] == account_key)
        self.assertIsNone(account["cells"][position]["value"], "only the ledger's items carry a balance")
        # No ledger on the page, no Balance column.
        plain = self.env["odin.pnl.export"]._payload(self.options(unfolded=["L:REV", account_key]))
        self.assertNotIn("Balance", [column["label"] for column in plain["columns"]])

    def test_the_payload_marks_notes_as_footnotes(self):
        self.report.add_annotation(self.options(), "L:REV", "One client")
        payload = self.env["odin.pnl.export"]._payload(self.options())
        self.assertEqual(payload["notes"][0]["text"], "One client")
        revenue = next(row for row in payload["rows"] if row["key"] == "L:REV")
        self.assertEqual(revenue["note_refs"], [1])

    def test_the_pdf_renders(self):
        if self.env["ir.actions.report"].get_wkhtmltopdf_state() != "ok":
            self.skipTest("wkhtmltopdf is not installed")
        # Tests render reports as HTML unless asked otherwise (ir_actions_report.py).
        # wkhtmltopdf fetches the report's styles from this test server, which
        # holds the registry lock while tests run: allow_pdf_render (odoo.tests)
        # lets those requests through instead of deadlocking.
        export = self.env["odin.pnl.export"].with_context(force_report_rendering=True)
        with self.allow_pdf_render():
            content, filename, mimetype = export.render(self.options(), "pdf")
        self.assertEqual(mimetype, "application/pdf")
        self.assertTrue(content.startswith(b"%PDF"))
        html = self.env["ir.actions.report"]._render_qweb_html(
            "odin_account_pnl.action_report_pnl", None,
            data={"payload": self.env["odin.pnl.export"]._payload(self.options())})[0]
        self.assertIn(b"Net Profit", html)


@tagged("post_install", "-at_install")
class TestDetailLevels(PnlCase):
    """Lines, every account, every transaction: on screen and in exports,
    the levels of detail accountants expect of a P&L."""

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        for day, amount in (("2026-09-03", -500.0), ("2026-09-12", -300.0), ("2026-09-25", -200.0)):
            cls.post(day, [(cls.revenue, amount, {"partner": cls.partner_a, "label": f"Sale {day}"})])
        cls.post("2026-09-15", [(cls.opex, 150.0, {"label": "Rent"})])
        cls.account_key = f"L:REV/A:{cls.revenue.id}"

    def rows(self, **values):
        return self.report.get_report(self.options(**values))["rows"]

    def test_each_level_shows_what_it_says(self):
        types = {level: {row["type"] for row in self.rows(detail=level)} for level in ("lines", "accounts", "transactions")}
        self.assertNotIn("account", types["lines"])
        self.assertIn("account", types["accounts"])
        self.assertNotIn("entry", types["accounts"])
        self.assertIn("entry", types["transactions"])
        entries = [row for row in self.rows(detail="transactions") if row["type"] == "entry"]
        self.assertEqual(sorted(row["amount"] for row in entries), [150.0, 200.0, 300.0, 500.0])

    def test_views_saved_with_unfold_all_open_on_the_accounts(self):
        self.assertEqual(self.report.get_report(self.options(unfold_all=True))["options"]["detail"], "accounts")

    def test_the_screen_pages_the_items_but_an_export_lists_them_all(self):
        with patch.object(odin_pnl_report, "PAGE_SIZE", 2):
            screen = self.rows(detail="transactions")
            self.assertIn("more", {row["type"] for row in screen}, "two items, then Load more")
            payload = self.env["odin.pnl.export"]._payload(self.options(detail="transactions"))
        items = [row for row in payload["rows"] if row["type"] == "entry" and row["key"].startswith(self.account_key)]
        self.assertEqual(len(items), 3)
        self.assertFalse(payload["warnings"])

    def test_an_export_cut_short_says_so(self):
        with patch.object(odin_pnl_report, "EXPORT_ITEMS", 1):
            payload = self.env["odin.pnl.export"]._payload(self.options(detail="transactions"))
        self.assertTrue(any("Not every journal item" in warning for warning in payload["warnings"]))

    def test_an_unfolded_row_prints_its_figures_on_a_total_row(self):
        payload = self.env["odin.pnl.export"]._payload(self.options(detail="transactions"))
        statement = payload["statement"]
        header = next(index for index, row in enumerate(statement) if row["key"] == self.account_key)
        total = next(index for index, row in enumerate(statement) if row["key"] == f"{self.account_key}#total")
        self.assertTrue(all(cell["value"] is None for cell in statement[header]["cells"]), "a heading")
        self.assertEqual(statement[total]["label"], f"Total {self.revenue.name}")
        self.assertEqual(statement[total]["cells"][0]["value"], 1000.0)
        self.assertEqual([row["type"] for row in statement[header + 1:total]], ["entry"] * 3, "its items between")
        self.assertEqual(statement[total + 1]["label"], "Total Revenue", "then the line's own total")
        # The email reads the rows, where the line keeps its figure.
        revenue = next(row for row in payload["rows"] if row["key"] == "L:REV")
        self.assertEqual(revenue["cells"][0]["value"], 1000.0)

    def test_the_workbook_folds_and_lists_the_items_flat(self):
        content = self.env["odin.pnl.export"].render(self.options(detail="transactions"), "xlsx")[0]
        book = openpyxl.load_workbook(io.BytesIO(content))
        statement, items = book.worksheets
        header = [cell.value for cell in statement[5]]
        self.assertEqual(header[1:4], ["Date", "Contact", "Description"])
        rent = next(row for row in statement.iter_rows(min_row=6) if row[3].value == "Rent")
        self.assertGreater(statement.row_dimensions[rent[0].row].outline_level, 1, "folds under its account")
        self.assertEqual(rent[4].value, 150.0, "a number, not text")
        self.assertEqual(items.title, "Transactions")
        rows = [[cell.value for cell in row] for row in items.iter_rows(min_row=2)]
        self.assertEqual(len(rows), 4)
        sale = next(row for row in rows if row[7] == "Sale 2026-09-12")
        self.assertEqual(sale[0], "Revenue")
        self.assertEqual(sale[2], self.revenue.name)
        self.assertEqual(sale[6], "partner_a")
        self.assertEqual((sale[8], sale[9]), (300.0, -300.0), "as on the P&L, and debit minus credit")
